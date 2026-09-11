#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown → Word（适配 Word 的中文文档排版）· 科生 M2 报告交付用

与"直接 md 塞进 word"的区别（本工具的核心价值）：
  · 页面：A4 + 规范页边距；页眉（文档名）+ 页脚（第 X 页 / 共 Y 页，真域代码）
  · 封面页（标题/副标题/出品/日期）+ **自动目录**（TOC 域，打开自动更新）
  · 中文字体成套：标题 微软雅黑（加粗）/ 正文 宋体（西文 Times New Roman 回退）
  · 正文中文段首**首行缩进 2 字符**、行距 1.5；标题层级 1-4 级样式
  · **真 Word 表格**（Table Grid + 表头底纹），不是 "| a | b |" 文本
  · 内联格式：**加粗** / *斜体* / `等宽代码` / ~~删除线~~ / [链接](url)
  · 列表（项目符号 / 编号，支持多级缩进）；引用块（缩进+底纹）；代码块（等宽+底纹）
  · 图片 ![alt](path) 实际嵌入并按页宽等比缩放；分隔线 → 段落边框

用法:
  python -X utf8 scripts/md_to_docx.py <in.md> <out.docx> [选项]
选项:
  --title T        封面主标题（默认取 md 的一级标题）
  --subtitle S     封面副标题
  --author A       封面出品/机构（多行用 / 分隔）
  --date D         封面日期（默认当天）
  --no-cover       不生成封面
  --no-toc         不生成目录
  --header H       页眉文字（默认=标题）
"""
import argparse
import datetime
import io
import os
import re
import sys

# ⚠️ 不要在此处重绑 sys.stdout —— 本模块会被 build_science_report.py import，
# 模块级重绑会夺走 buffer 所有权导致调用方 "I/O operation on closed file"。改为在 main() 内处理。

from docx import Document                                    # noqa: E402
from docx.enum.section import WD_SECTION                     # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT               # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # noqa: E402
from docx.oxml import OxmlElement                            # noqa: E402
from docx.oxml.ns import qn                                  # noqa: E402
from docx.shared import Cm, Pt, RGBColor                     # noqa: E402

# ---------- 字体（成套，可按品牌 VI 覆盖）----------
FONT_TITLE = "微软雅黑"     # 标题
FONT_BODY = "宋体"          # 正文中文
FONT_LATIN = "Times New Roman"  # 正文西文
FONT_MONO = "Consolas"      # 等宽
SIZE_BODY = Pt(10.5)        # 五号
LINE = 1.5                  # 行距
INDENT_CHARS = 2            # 首行缩进字符数


# ==================== 底层工具 ====================
def _ea(el, font):
    """给某元素（run/style 的 rPr）设置中日韩字体。"""
    rPr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else el
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font if font != FONT_BODY else FONT_LATIN)
    rFonts.set(qn("w:hAnsi"), font if font != FONT_BODY else FONT_LATIN)
    rFonts.set(qn("w:eastAsia"), font)


def set_run_font(run, font=None, size=None, bold=None, italic=None, color=None, strike=None):
    if font:
        run.font.name = font
        _ea(run._element, font)
    if size:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if strike is not None:
        run.font.strike = strike
    if color:
        run.font.color.rgb = RGBColor(*color)


def _shade(el, fill):
    """给段落/单元格底纹。"""
    pr = el.get_or_add_pPr() if el.tag.endswith("}p") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _p_border(p, edge="bottom", sz=6, color="999999"):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "1")
    e.set(qn("w:color"), color)
    bdr.append(e)


def add_field(paragraph, instr, placeholder="…", font=None, size=None, color=None):
    """在段落里插入 Word 域代码（如 TOC / PAGE / NUMPAGES）。"""
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for x in (b, i, s, t, e):
        run._r.append(x)
    if font or size or color:
        set_run_font(run, font, size, color=color)
    return run


def update_fields_on_open(doc):
    """让 Word 打开时自动更新目录/页码等域。"""
    st = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    st.append(uf)


# ==================== 文档骨架 ====================
def init_doc():
    doc = Document()
    # 页面 A4 + 页边距
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.8)
        sec.top_margin = sec.bottom_margin = Cm(2.6)
    # 正文样式
    st = doc.styles["Normal"]
    st.font.size = SIZE_BODY
    st.font.name = FONT_LATIN
    _ea(st.element, FONT_BODY)
    pf = st.paragraph_format
    pf.line_spacing = LINE
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(SIZE_BODY.pt * INDENT_CHARS)   # 首行缩进 2 字符
    # 标题样式（中文字体成套）
    for lvl, (sz, col) in enumerate([(Pt(18), (0x1F, 0x38, 0x64)), (Pt(15), (0x1F, 0x38, 0x64)),
                                     (Pt(13), (0x2E, 0x54, 0x96)), (Pt(12), (0x2E, 0x54, 0x96))], 1):
        s = doc.styles[f"Heading {lvl}"]
        s.font.size = sz
        s.font.bold = True
        s.font.name = FONT_LATIN
        s.font.color.rgb = RGBColor(*col)
        _ea(s.element, FONT_TITLE)
        s.paragraph_format.first_line_indent = Pt(0)
        s.paragraph_format.space_before = Pt(12 if lvl <= 2 else 8)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.line_spacing = 1.25
    # 列表样式去缩进残留
    for nm in ("List Bullet", "List Number", "List Bullet 2", "List Number 2"):
        try:
            doc.styles[nm].paragraph_format.first_line_indent = Pt(0)
        except KeyError:
            pass
    return doc


def add_cover(doc, title, subtitle, author, date):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(140)
    p.paragraph_format.first_line_indent = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title or "报告"); set_run_font(r, FONT_TITLE, Pt(30), bold=True, color=(0x1F, 0x38, 0x64))
    if subtitle:
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.first_line_indent = Pt(0)
        r2 = p2.add_run(subtitle); set_run_font(r2, FONT_TITLE, Pt(14), color=(0x55, 0x55, 0x55))
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(90); p3.paragraph_format.first_line_indent = Pt(0)
    r3 = p3.add_run(f"{author}\n{date}" if author else date)
    set_run_font(r3, FONT_BODY, Pt(12), color=(0x66, 0x66, 0x66))
    doc.add_page_break()


def add_toc(doc, levels="1-3"):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("目录"); set_run_font(r, FONT_TITLE, Pt(16), bold=True, color=(0x1F, 0x38, 0x64))
    p2 = doc.add_paragraph(); p2.paragraph_format.first_line_indent = Pt(0)
    add_field(p2, f'TOC \\o "{levels}" \\h \\z \\u', "（目录：打开文档后按 F9 或右键→更新域）")
    doc.add_page_break()


def add_header_footer(doc, header_text):
    for sec in doc.sections:
        # 页眉
        hp = sec.header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr = hp.add_run(header_text or "")
        set_run_font(hr, FONT_BODY, Pt(9), color=(0x88, 0x88, 0x88))
        _p_border(hp, "bottom")
        # 页脚：第 X 页 / 共 Y 页
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = fp.add_run("第 "); set_run_font(r1, FONT_BODY, Pt(9), color=(0x88, 0x88, 0x88))
        add_field(fp, "PAGE", "1", FONT_BODY, Pt(9), (0x88, 0x88, 0x88))
        r2 = fp.add_run(" 页 / 共 "); set_run_font(r2, FONT_BODY, Pt(9), color=(0x88, 0x88, 0x88))
        add_field(fp, "NUMPAGES", "1", FONT_BODY, Pt(9), (0x88, 0x88, 0x88))
        r3 = fp.add_run(" 页"); set_run_font(r3, FONT_BODY, Pt(9), color=(0x88, 0x88, 0x88))


# ==================== 内联解析 ====================
TOKEN = re.compile(
    r"(\*\*.+?\*\*|__.+?__|\*[^*\n]+?\*|_[^_\n]+?_|`[^`\n]+?`|~~.+?~~|\[[^\]]+?\]\([^)]+?\))"
)


def add_inline(paragraph, text, base_font=FONT_BODY, base_size=SIZE_BODY, bold_all=False):
    """把一行 markdown 内联语法渲染成 runs。"""
    for seg in TOKEN.split(text):
        if not seg:
            continue
        b, i, code, strike = bold_all, False, False, False
        s = seg
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            s, b = seg[2:-2], True
        elif seg.startswith("__") and seg.endswith("__") and len(seg) > 4:
            s, b = seg[2:-2], True
        elif seg.startswith("~~") and seg.endswith("~~") and len(seg) > 4:
            s, strike = seg[2:-2], True
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            s, code = seg[1:-1], True
        elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            s, i = seg[1:-1], True
        elif seg.startswith("_") and seg.endswith("_") and len(seg) > 2:
            s, i = seg[1:-1], True
        elif seg.startswith("[") and "](" in seg:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", seg)
            if m:
                s = f"{m.group(1)}（{m.group(2)}）"
        r = paragraph.add_run(s)
        set_run_font(r, FONT_MONO if code else base_font,
                     Pt(base_size.pt - 0.5) if code else base_size,
                     bold=b, italic=i, strike=strike,
                     color=(0xC0, 0x39, 0x2B) if code else None)
    return paragraph


# ==================== 块解析 ====================
def parse_blocks(lines):
    """把 md 行切成块：heading/table/list/quote/code/hr/image/para。"""
    blocks, i, n = [], 0, len(lines)
    while i < n:
        ln = lines[i]
        s = ln.rstrip()
        if not s.strip():
            i += 1; continue
        # 代码块
        if s.lstrip().startswith("```"):
            lang = s.lstrip()[3:].strip()
            buf, i = [], i + 1
            while i < n and not lines[i].lstrip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            blocks.append(("code", "\n".join(buf), lang)); continue
        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("h", m.group(2).strip(), len(m.group(1)))); i += 1; continue
        # 分隔线
        if re.match(r"^\s*([-*_])\s*(\1\s*){2,}$", s):
            blocks.append(("hr", "", 0)); i += 1; continue
        # 表格（连续 | 行）
        if s.lstrip().startswith("|"):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            blocks.append(("table", rows, 0)); continue
        # 引用块
        if s.lstrip().startswith(">"):
            buf = []
            while i < n and lines[i].lstrip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            blocks.append(("quote", "\n".join(buf), 0)); continue
        # 列表
        m = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", ln)
        if m:
            items = []
            while i < n:
                mm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", lines[i])
                if not mm:
                    break
                indent = len(mm.group(1)) // 2
                ordered = bool(re.match(r"\d", mm.group(2)))
                items.append((indent, ordered, mm.group(3).strip()))
                i += 1
            blocks.append(("list", items, 0)); continue
        # 图片（独占一行）
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", s)
        if m:
            blocks.append(("img", m.group(2).strip(), m.group(1).strip())); i += 1; continue
        # 普通段落（合并到空行）
        buf = [s]
        i += 1
        while i < n and lines[i].strip() and not re.match(
                r"^(#{1,6}\s|\s*[-*+]\s|\s*\d+[.)]\s|\s*>|\s*\||\s*```|\s*!\[)", lines[i]):
            buf.append(lines[i].rstrip()); i += 1
        blocks.append(("para", "\n".join(buf), 0))
    return blocks


def render_table(doc, rows):
    # 去分隔行 |---|---|
    data = []
    for idx, r in enumerate(rows):
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if idx == 1 and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != ""):
            continue
        data.append(cells)
    if not data:
        return
    ncol = max(len(r) for r in data)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(data):
        cells = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, txt, base_size=Pt(10), bold_all=(ri == 0))
            if ri == 0:
                _shade(cell._tc, "DCE6F1")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render(doc, blocks, base_dir):
    for kind, payload, meta in blocks:
        if kind == "h":
            lvl = min(meta, 4)
            doc.add_heading(payload, level=lvl)
        elif kind == "para":
            for seg in payload.split("\n"):
                if not seg.strip():
                    continue
                p = doc.add_paragraph()
                add_inline(p, seg.strip())
        elif kind == "list":
            for indent, ordered, txt in payload:
                style = ("List Number" if ordered else "List Bullet") + (" 2" if indent else "")
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:
                    p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
                p.paragraph_format.first_line_indent = Pt(0)
                add_inline(p, txt)
        elif kind == "quote":
            for seg in payload.split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                _shade(p._p, "F2F4F7")
                add_inline(p, seg.strip(), base_size=Pt(10), base_font=FONT_BODY)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            _shade(p._p, "F5F5F5")
            r = p.add_run(payload)
            set_run_font(r, FONT_MONO, Pt(9.5), color=(0x33, 0x33, 0x33))
        elif kind == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            _p_border(p, "bottom", sz=8)
        elif kind == "img":
            path = payload
            if not os.path.isabs(path):
                path = os.path.normpath(os.path.join(base_dir, path))
            if os.path.exists(path):
                try:
                    doc.add_picture(path, width=Cm(15.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if meta:
                        cap = doc.add_paragraph()
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap.paragraph_format.first_line_indent = Pt(0)
                        cr = cap.add_run(f"图：{meta}")
                        set_run_font(cr, FONT_BODY, Pt(9), color=(0x80, 0x80, 0x80))
                except Exception as e:  # noqa: BLE001
                    p = doc.add_paragraph()
                    add_inline(p, f"[图片插入失败：{os.path.basename(path)} — {e}]")
            else:
                p = doc.add_paragraph()
                add_inline(p, f"[图片缺失：{path}]")
        elif kind == "table":
            render_table(doc, payload)


def build(md_path, out_path, title=None, subtitle=None, author=None, date=None,
          cover=True, toc=True, header=None):
    with io.open(md_path, encoding="utf-8-sig") as f:
        lines = f.read().split("\n")
    blocks = parse_blocks(lines)
    doc = init_doc()
    # 一级标题作为文档题名（正文里不再重复渲染一次大标题）
    h1 = next((p for k, p, _m in blocks if k == "h" and _m == 1), None)
    doc_title = title or h1 or os.path.splitext(os.path.basename(md_path))[0]
    if cover:
        add_cover(doc, doc_title, subtitle, author, date or datetime.date.today().isoformat())
    if toc:
        add_toc(doc)
    add_header_footer(doc, header or doc_title)
    # 去掉首个大标题（已作封面/页眉题名）
    if h1 and not title:
        for idx, (k, p, m) in enumerate(blocks):
            if k == "h" and m == 1:
                blocks.pop(idx); break
    render(doc, blocks, os.path.dirname(os.path.abspath(md_path)))
    update_fields_on_open(doc)
    doc.save(out_path)
    return doc_title


def main():
    try:
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass
    ap = argparse.ArgumentParser(description="Markdown → Word（适配 Word 的中文排版）")
    ap.add_argument("md"); ap.add_argument("out")
    ap.add_argument("--title"); ap.add_argument("--subtitle"); ap.add_argument("--author")
    ap.add_argument("--date"); ap.add_argument("--header")
    ap.add_argument("--no-cover", action="store_true"); ap.add_argument("--no-toc", action="store_true")
    a = ap.parse_args()
    if not os.path.exists(a.md):
        print(f"输入不存在: {a.md}"); sys.exit(2)
    t = build(a.md, a.out, a.title, a.subtitle, a.author, a.date,
              not a.no_cover, not a.no_toc, a.header)
    print(f"== 已生成 Word（{t}）-> {a.out}")
    print("   含：封面" + ("/目录" if not a.no_toc else "") + "/页眉页脚页码/中文字体/首行缩进/真表格")


if __name__ == "__main__":
    main()
